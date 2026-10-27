"""Export twenty report pages from the actual AURA source, without rewriting it.
Run with a Python environment containing reportlab, python-docx, and pypdf.
"""
from pathlib import Path
import hashlib
import json
import textwrap
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)
# Every range is inclusive and refers to the unmodified source file.
PAGES = [
    ("Application startup", "React frontend entry point and FastAPI service initialization.", [
        ("frontend/src/main.jsx", 1, 18), ("backend/app/main.py", 1, 46)]),
    ("Authentication and access control", "Password hashing, token creation and authenticated-user lookup.", [
        ("backend/app/auth.py", 1, 58)]),
    ("Database and session persistence", "SQLite connection, database sessions and ORM user/work-session models.", [
        ("backend/app/database.py", 1, 35), ("backend/app/models.py", 9, 40)]),
    ("Feature schema and API communication", "Validated multimodal inputs, browser API calls and prediction endpoint entry.", [
        ("backend/app/schemas.py", 47, 70), ("frontend/src/api.js", 16, 36),
        ("backend/app/routers/fatigue.py", 48, 58)]),
    ("Face model loading and health checks", "Local model/WASM loading, CPU inference, cancellation and freshness checks.", [
        ("frontend/src/lib/faceEngine.js", 1, 62)]),
    ("Camera startup and tracking watchdog", "Owned engine startup, camera acquisition and stale-tracking cleanup.", [
        ("frontend/src/hooks/useFaceMesh.js", 195, 229),
        ("frontend/src/hooks/useFaceMesh.js", 349, 364)]),
    ("Real-time facial feature extraction", "Fresh video-frame inference, eye/mouth features, pose and action tracking.", [
        ("frontend/src/hooks/useFaceMesh.js", 249, 288),
        ("frontend/src/hooks/useFaceMesh.js", 297, 322)]),
    ("Eye and mouth landmark geometry", "EAR, per-eye closure, corrected lip pairs, MAR and matrix-based head pose.", [
        ("frontend/src/lib/vision.js", 1, 59)]),
    ("Conservative yawn filtering", "Mouth/jaw agreement, speech-like movement suppression and closure confirmation.", [
        ("frontend/src/lib/yawn.js", 62, 101), ("frontend/src/lib/yawn.js", 120, 135)]),
    ("Typing and mouse measurement", "Single-keydown counting, pointer distance/time and focused input collection.", [
        ("extension/activity-sampler.js", 1, None),
        ("frontend/src/hooks/useActivity.js", 20, 36)]),
    ("Personal starting-pace baseline", "First active-minute reference, sustained slowdown and bounded score adjustment.", [
        ("frontend/src/lib/activity.js", 1, 47), ("frontend/src/lib/activity.js", 54, 64),
        ("frontend/src/lib/activity.js", 71, 75)]),
    ("LightGBM training and fatigue scoring", "Synthetic-data training, prediction and the explicit personal-pace adjustment.", [
        ("backend/app/ml/train.py", 108, 142),
        ("backend/app/ml/fatigue_model.py", 160, 178),
        ("backend/app/ml/activity.py", 1, 16)]),
    ("LinUCB intervention personalization", "Upper-confidence selection, reward updates and break-trigger gating.", [
        ("backend/app/ml/bandit.py", 89, 103), ("backend/app/ml/bandit.py", 123, 163),
        ("backend/app/ml/bandit.py", 192, 197)]),
    ("Selectable beep sounds and preview", "Saved sound/volume preferences, preview playback, fatigue mute and away tone.", [
        ("frontend/src/lib/sound.js", 11, 32), ("frontend/src/lib/sound.js", 53, 89)]),
    ("Fatigue alerts and reminder scheduling", "Sustained score-75 warnings and persisted Guardian scheduled notifications.", [
        ("frontend/src/pages/Monitor.jsx", 212, 241),
        ("extension/background.js", 204, 239)]),
]


PAGES.extend([
    ("Sound-picker user interface", "React sound selection, saved preferences, preview actions and accessible status.", [
        ("frontend/src/components/SoundPicker.jsx", 1, 7),
        ("frontend/src/components/SoundPicker.jsx", 21, 41),
        ("frontend/src/components/SoundPicker.jsx", 44, 62)]),
    ("Browser fallback fatigue scoring", "The rule-based fallback used by the public static app, separate from Python ML.", [
        ("frontend/src/lib/localEngine.js", 79, 116)]),
    ("Ember Docket persistence and delivery", "Local task storage, due-time checks and Guardian delivery acknowledgements.", [
        ("frontend/src/lib/docket.js", 1, 65)]),
    ("Guardian background audio engine", "Persisted sound settings, offscreen playback, looping away alarm and manual stop.", [
        ("extension/background.js", 29, 47), ("extension/offscreen.js", 1, 38)]),
    ("Session summary and feedback learning", "Stored session statistics, user feedback rewards and personalized bandit updates.", [
        ("backend/app/routers/fatigue.py", 23, 45),
        ("backend/app/routers/fatigue.py", 135, 163)]),
])
PAGE_COUNT = len(PAGES)
assert PAGE_COUNT == 20

FONT_DIR = Path("/usr/share/fonts/truetype/dejavu")
for name, filename in [("CodeMono", "DejaVuSansMono.ttf"), ("Body", "DejaVuSans.ttf"), ("Heading", "DejaVuSans-Bold.ttf")]:
    pdfmetrics.registerFont(TTFont(name, str(FONT_DIR / filename)))

blocks = []
manifest = []
for title, note, ranges in PAGES:
    selected = []
    for rel, start, end in ranges:
        data = (ROOT / rel).read_text()
        lines = data.splitlines()
        end = end or len(lines)
        assert 0 < start <= end <= len(lines), (rel, start, end)
        code = "\n".join(lines[start - 1:end])
        selected.append((rel, start, end, code))
        manifest.append({"file": rel, "start": start, "end": end,
                         "sha256": hashlib.sha256(code.encode()).hexdigest()})
    blocks.append((title, note, selected))

pdf_path = OUT / "AURA_20_Page_Project_Code.pdf"
docx_path = OUT / "AURA_20_Page_Project_Code.docx"
md_path = OUT / "AURA_20_Page_Project_Code.md"
W, H = A4
LEFT = 36
RIGHT = W - 36
WIDTH = RIGHT - LEFT
pdf = canvas.Canvas(str(pdf_path), pagesize=A4)
pdf.setTitle("AURA - 20-page project source-code listings")
pdf.setAuthor("AURA project")

doc = Document()
sec = doc.sections[0]
sec.page_width = Pt(W); sec.page_height = Pt(H)
sec.top_margin = Pt(36); sec.bottom_margin = Pt(36)
sec.left_margin = Pt(36); sec.right_margin = Pt(36)
sec.header_distance = Pt(15); sec.footer_distance = Pt(15)
normal = doc.styles["Normal"]
normal.font.name = "DejaVu Sans"; normal.font.size = Pt(9)
normal.paragraph_format.space_after = Pt(0)
footer = sec.footer.paragraphs[0]
footer.alignment = 2
r = footer.add_run("AURA | Project source-code excerpts | Page ")
r.font.size = Pt(7)
field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
footer.add_run(" / 20").font.size = Pt(7)

md = ["# AURA — 20-page project code", "",
      "Actual source excerpts from the current project checkout, prepared 27 October 2026.",
      "These are report listings, not a standalone runnable replacement for the repository. Imports, enclosing scopes, UI, assets and other dependencies remain in the full project.",
      "The PDF is exactly 20 A4 pages. The Word file has matching explicit page breaks; pagination can vary with fonts/printer settings. Long PDF/Word lines wrap visually; use the unchanged Markdown code blocks for copying.",
      "Training data is synthetic, not clinical. Personal work-pace scoring is a bounded heuristic; the public static app uses browser fallback scoring, not the Python model server.", ""]
font_sizes = []
source_line_count = 0
for number, (title, note, selected) in enumerate(blocks, 1):
    # Reduce only the font slightly on the denser pages; never drop source lines.
    size = 8.4
    while True:
        max_chars = int(WIDTH / pdfmetrics.stringWidth("M", "CodeMono", size))
        rows = []
        for rel, start, end, code in selected:
            rows.append(("label", f"{rel} | lines {start}-{end}"))
            for line in code.split("\n"):
                wrapped = textwrap.wrap(line.expandtabs(4), width=max_chars, replace_whitespace=False,
                                        drop_whitespace=False, break_long_words=True, break_on_hyphens=False) or [""]
                rows.extend(("code", piece) for piece in wrapped)
            rows.append(("gap", ""))
        leading = size + 1.5
        required = sum(14 if kind == "label" else 6 if kind == "gap" else leading for kind, _ in rows)
        top = H - (125 if number == 1 else 100)
        if required <= top - 47:
            break
        size = round(size - .1, 1)
        assert size >= 7.0, (number, required)
    font_sizes.append(size)
    pdf.setFillColor(HexColor("#173d33")); pdf.rect(0, H - 10, W, 10, fill=1, stroke=0)
    pdf.setFillColor(HexColor("#58706a")); pdf.setFont("Heading", 8)
    pdf.drawString(LEFT, H - 34, "AURA / MINI-PROJECT IMPLEMENTATION")
    pdf.setFillColor(HexColor("#102a23")); pdf.setFont("Heading", 15)
    pdf.drawString(LEFT, H - 58, f"{number:02d}. {title}")
    pdf.setFillColor(HexColor("#42554f")); pdf.setFont("Body", 8)
    pdf.drawString(LEFT, H - 77, note)
    if number == 1:
        pdf.setFont("Body", 7.7)
        pdf.drawString(LEFT, H - 94, "Actual source excerpts for your report; not a standalone runnable project.")
        pdf.drawString(LEFT, H - 105, "Long lines wrap visually. Use the Markdown companion for exact, copyable code.")
    y = top
    for kind, text in rows:
        if kind == "label":
            pdf.setFillColor(HexColor("#e9f1ed")); pdf.rect(LEFT-3, y-3, WIDTH+6, 12, fill=1, stroke=0)
            pdf.setFillColor(HexColor("#214d3e")); pdf.setFont("Body", 7.5); pdf.drawString(LEFT, y, text); y -= 14
        elif kind == "gap":
            y -= 6
        else:
            pdf.setFillColor(HexColor("#19231f")); pdf.setFont("CodeMono", size); pdf.drawString(LEFT, y, text); y -= leading
    assert y >= 42, (number, y)
    pdf.setStrokeColor(HexColor("#cad8d0")); pdf.line(LEFT, 33, RIGHT, 33)
    pdf.setFillColor(HexColor("#597168")); pdf.setFont("Body", 7)
    pdf.drawString(LEFT, 21, "Source excerpts | Dependencies and omitted context remain in the full project")
    pdf.drawRightString(RIGHT, 21, f"{number:02d} / 20")
    pdf.showPage()

    if number > 1:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(f"AURA / MINI-PROJECT IMPLEMENTATION\n{number:02d}. {title}")
    run.bold = True; run.font.size = Pt(13); run.font.color.rgb = RGBColor.from_string("173D33")
    p = doc.add_paragraph(note); p.paragraph_format.space_after = Pt(6)
    p.runs[0].font.size = Pt(8)
    if number == 1:
        p = doc.add_paragraph("Actual source excerpts, not a standalone runnable project. Use the Markdown companion to copy exact code; dependencies and omitted context remain in the full repository.")
        p.paragraph_format.space_after = Pt(6); p.runs[0].font.size = Pt(8)
    # Same explicit visual line wrapping as the PDF, no auto-numbering or smart quotes.
    for kind, text in rows:
        p = doc.add_paragraph()
        fmt = p.paragraph_format
        fmt.space_before = Pt(0); fmt.space_after = Pt(0)
        fmt.line_spacing = Pt(14 if kind == "label" else 6 if kind == "gap" else leading)
        fmt.keep_with_next = kind == "label"
        fmt.widow_control = False
        run = p.add_run(text)
        run.font.name = "DejaVu Sans Mono" if kind == "code" else "DejaVu Sans"
        run.font.size = Pt(size if kind == "code" else 7.5)
        if kind == "label":
            run.bold = True; run.font.color.rgb = RGBColor.from_string("214D3E")
    md.extend([f"## Page {number:02d} — {title}", "", note, ""])
    for rel, start, end, code in selected:
        lang = "python" if rel.endswith(".py") else "jsx" if rel.endswith(".jsx") else "javascript"
        md.extend([f"### `{rel}` — lines {start}–{end}", "", f"```{lang}", code, "```", ""])
        source_line_count += len(code.split("\n"))

pdf.save(); doc.save(str(docx_path)); md_path.write_text("\n".join(md), encoding="utf-8")
assert len(PdfReader(str(pdf_path)).pages) == PAGE_COUNT
for i, page in enumerate(PdfReader(str(pdf_path)).pages, 1):
    assert f"{i:02d} / 20" in page.extract_text(), i
# Re-read every listing to check it still matches its source byte-for-byte (UTF-8 code).
for item in manifest:
    exact = "\n".join((ROOT/item['file']).read_text().splitlines()[item['start']-1:item['end']])
    assert hashlib.sha256(exact.encode()).hexdigest() == item['sha256']
print(json.dumps({"pdf_pages": PAGE_COUNT, "source_lines": source_line_count, "listings": len(manifest),
                  "code_font_sizes_pt": font_sizes, "files": [str(pdf_path), str(docx_path), str(md_path)]}, indent=2))

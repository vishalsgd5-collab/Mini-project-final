"""Export five full-page A4 reviews, preserving one body paragraph per review.

Requires reportlab, python-docx and pypdf. Run from any working directory.
The validated PDF is the fixed-layout version; Word may reflow with font changes.
"""
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables'
SOURCE = OUT / 'AURA_Literature_Review_Full_Page_Content.md'
PDF = OUT / 'AURA_Literature_Review_5_Full_Pages.pdf'
DOCX = OUT / 'AURA_Literature_Review_5_Full_Pages.docx'
sections = []
for block in SOURCE.read_text(encoding='utf-8').split('\n### ')[1:]:
    heading, text = block.split('\n', 1)
    text = text.strip()
    assert '\n\n' not in text, 'Keep one continuous paragraph per review.'
    sections.append((heading, text))
assert len(sections) == 5

# Embed a Unicode font for author names such as SOUKUPOVÁ and ČECH.
pdfmetrics.registerFont(TTFont('ReviewHeading', '/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf'))
W, H = A4
MARGIN = 50
WIDTH = W - 2*MARGIN
BOTTOM = 51
HEADING_TOP = H - 70
heading_style = ParagraphStyle('Paper', fontName='ReviewHeading', fontSize=10.7,
                               leading=14.2, textColor=HexColor('#173f36'))
pdf = canvas.Canvas(str(PDF), pagesize=A4)
pdf.setTitle('AURA - Literature Review - Five Full A4 Pages')
pdf.setAuthor('AURA project report')
doc = Document()
section = doc.sections[0]
section.page_width = Pt(W); section.page_height = Pt(H)
section.left_margin = Pt(MARGIN); section.right_margin = Pt(MARGIN)
section.top_margin = Pt(30); section.bottom_margin = Pt(46)
section.footer_distance = Pt(20)
style = doc.styles['Normal']
style.font.name = 'Times New Roman'; style.font.size = Pt(12)
style.paragraph_format.space_after = Pt(0)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('AURA | Literature Review | Page ').font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
footer.add_run(' of 5').font.size = Pt(8)

for number, (heading, text) in enumerate(sections, 1):
    title = Paragraph(escape(heading), heading_style)
    _, title_height = title.wrap(WIDTH, 10000)
    body_top = HEADING_TOP - title_height - 12
    size = 12
    while True:
        probe = ParagraphStyle('Measure', fontName='Times-Roman', fontSize=size, leading=1)
        lines = Paragraph(escape(text), probe).wrap(WIDTH, 10000)[1]
        leading = (body_top - BOTTOM)/lines
        if leading >= 14: break
        size = round(size-.1, 1)
        assert size >= 11.5, (number, leading)
    assert leading <= 18, 'Text too short for a full page at the intended density.'
    body_style = ParagraphStyle('Body', fontName='Times-Roman', fontSize=size,
                                leading=leading, alignment=TA_JUSTIFY,
                                textColor=HexColor('#202828'))
    body = Paragraph(escape(text), body_style)
    _, body_height = body.wrap(WIDTH, 10000)
    assert abs(body_top-body_height-BOTTOM)<.1
    pdf.setFillColor(HexColor('#173f36')); pdf.rect(0, H-10, W, 10, stroke=0, fill=1)
    pdf.setFont('Helvetica-Bold', 11)
    pdf.drawString(MARGIN, H-37, 'LITERATURE REVIEW')
    pdf.setFont('Helvetica', 9)
    pdf.drawString(MARGIN, H-54, '2.1 LITERATURE REVIEW')
    title.drawOn(pdf, MARGIN, HEADING_TOP-title_height)
    body.drawOn(pdf, MARGIN, body_top-body_height)
    pdf.setStrokeColor(HexColor('#c8d8ce')); pdf.line(MARGIN, 35, W-MARGIN, 35)
    pdf.setFont('Helvetica', 7.4); pdf.setFillColor(HexColor('#506c63'))
    pdf.drawString(MARGIN, 22, 'AURA / AI-Based Fatigue Detection and Wellness Care System')
    pdf.drawRightString(W-MARGIN, 22, f'{number} / 5')
    pdf.showPage()

    if number > 1: doc.add_page_break()
    p = doc.add_paragraph('LITERATURE REVIEW')
    p.paragraph_format.space_after = Pt(5)
    p.runs[0].bold = True; p.runs[0].font.size = Pt(11)
    p = doc.add_paragraph('2.1 LITERATURE REVIEW')
    p.paragraph_format.space_after = Pt(9)
    p.runs[0].font.size = Pt(9)
    p = doc.add_paragraph(heading)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(14.2)
    p.paragraph_format.space_after = Pt(10)
    p.runs[0].bold = True; p.runs[0].font.size = Pt(11)
    p.runs[0].font.color.rgb = RGBColor.from_string('173F36')
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.widow_control = False
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(leading-.35)
    p.runs[0].font.size = Pt(size)
    print(f'Page {number}: {len(text.split())} words; {size}pt; leading {leading:.2f}pt')

pdf.save(); doc.save(DOCX)
reader = PdfReader(PDF)
assert len(reader.pages) == 5
for index, ((heading, text), page) in enumerate(zip(sections, reader.pages), 1):
    extracted = ' '.join(page.extract_text().split())
    assert f'2.1.{index}' in extracted and f'{index} / 5' in extracted
    assert ' '.join(text.split()) in extracted, f'Missing body text on page {index}'
    assert tuple(round(float(v),1) for v in page.mediabox[2:]) == (595.3,841.9)
print('Validated: exactly five A4 pages, each containing its complete review.')
print(PDF); print(DOCX)

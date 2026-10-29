"""Create one A4 report page per SDG from the companion Markdown content."""
from pathlib import Path
from xml.sax.saxutils import escape
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.colors import HexColor
from reportlab.platypus import Paragraph
from pypdf import PdfReader
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables'
source = OUT / 'AURA_SDG_Full_Page_Content.md'
sections = []
for block in source.read_text().split('\n## ')[1:]:
    title, rest = block.split('\n', 1)
    parts = rest.strip().split('\n\n')
    relation = parts[0].strip('*')
    sections.append((title, relation, parts[1:]))
assert len(sections) == 7

pdf_path = OUT / 'AURA_SDG_One_Full_Page_Per_Goal.pdf'
word_path = OUT / 'AURA_SDG_One_Full_Page_Per_Goal.docx'
W, H = A4
margin = 50
width = W - margin * 2
body_top = H - 112
body_bottom = 50
body_height = body_top - body_bottom
pdf = canvas.Canvas(str(pdf_path), pagesize=A4)
pdf.setTitle('AURA - Sustainable Development Goals - One Full A4 Page Per Goal')
pdf.setAuthor('AURA project report')
doc = Document()
section = doc.sections[0]
section.page_width = Pt(W); section.page_height = Pt(H)
section.left_margin = Pt(margin); section.right_margin = Pt(margin)
section.top_margin = Pt(40); section.bottom_margin = Pt(46)
section.footer_distance = Pt(20)
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'; normal.font.size = Pt(12)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('AURA | SDG alignment | Page ').font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
footer.add_run(' of 7').font.size = Pt(8)
measurements = []
for index, (title, relation, paragraphs) in enumerate(sections, 1):
    size = 12
    gap = 9
    while True:
        probe = ParagraphStyle('probe', fontName='Times-Roman', fontSize=size, leading=1, alignment=TA_JUSTIFY)
        line_count = sum(Paragraph(escape(text), probe).wrap(width, 10000)[1] for text in paragraphs)
        leading = min(18, (body_height - (len(paragraphs)-1)*gap) / line_count)
        if leading >= 14: break
        size = round(size-.1, 1)
        assert size >= 11.5, (index, line_count, leading)
    style = ParagraphStyle('body', fontName='Times-Roman', fontSize=size, leading=leading,
                           alignment=TA_JUSTIFY, textColor=HexColor('#202c28'))
    pdf.setFillColor(HexColor('#173f36')); pdf.rect(0, H-10, W, 10, fill=1, stroke=0)
    pdf.setFont('Helvetica-Bold', 8); pdf.setFillColor(HexColor('#506c63'))
    pdf.drawString(margin, H-36, 'AURA / SUSTAINABLE DEVELOPMENT GOALS')
    pdf.setFont('Times-Bold', 17)
    pdf.setFillColor(HexColor('#173f36'))
    pdf.drawString(margin, H-62, title)
    pdf.setFont('Helvetica', 9)
    pdf.drawString(margin, H-81, 'AI-Based Fatigue Detection and Wellness Care System')
    pdf.setFont('Helvetica-Oblique', 8.4)
    pdf.drawString(margin, H-100, relation)
    y = body_top
    for number, text in enumerate(paragraphs):
        paragraph = Paragraph(escape(text), style)
        _, height = paragraph.wrap(width, 10000)
        paragraph.drawOn(pdf, margin, y-height)
        y -= height
        if number < len(paragraphs)-1: y -= gap
    assert body_bottom-1 <= y <= body_bottom+35, (index, y)
    pdf.setStrokeColor(HexColor('#c8d8ce')); pdf.line(margin, 35, W-margin, 35)
    pdf.setFont('Helvetica', 7.4); pdf.setFillColor(HexColor('#506c63'))
    pdf.drawString(margin, 22, 'Project alignment and intended contributions; not verified SDG impact.')
    pdf.drawRightString(W-margin, 22, f'{index} / 7')
    pdf.showPage()

    if index > 1: doc.add_page_break()
    p = doc.add_paragraph('AURA / SUSTAINABLE DEVELOPMENT GOALS')
    p.paragraph_format.space_after = Pt(6)
    p.runs[0].font.name = 'Arial'; p.runs[0].font.size = Pt(8)
    p = doc.add_paragraph(title)
    p.paragraph_format.space_after = Pt(6)
    p.runs[0].bold = True; p.runs[0].font.size = Pt(17)
    p.runs[0].font.color.rgb = RGBColor.from_string('173F36')
    p = doc.add_paragraph('AI-Based Fatigue Detection and Wellness Care System')
    p.paragraph_format.space_after = Pt(5); p.runs[0].font.size = Pt(9)
    p = doc.add_paragraph(relation)
    p.paragraph_format.space_after = Pt(9); p.runs[0].italic = True; p.runs[0].font.size = Pt(8.5)
    for number, text in enumerate(paragraphs):
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(leading-.25)
        p.paragraph_format.space_after = Pt(gap if number < len(paragraphs)-1 else 0)
        p.paragraph_format.widow_control = False
        p.runs[0].font.name = 'Times New Roman'; p.runs[0].font.size = Pt(size)
    measurements.append((index, title, len(' '.join(paragraphs).split()), round(leading,2), round(y,2)))
pdf.save(); doc.save(word_path)
reader = PdfReader(pdf_path)
assert len(reader.pages) == 7
for number, (page, (title, _, paragraphs)) in enumerate(zip(reader.pages, sections),1):
    text = page.extract_text()
    assert title in text and f'{number} / 7' in text
    assert len(text.split()) >= 530
print('Verified 7 A4 pages with one complete SDG section per page.')
for row in measurements: print(row)
print(pdf_path); print(word_path)
